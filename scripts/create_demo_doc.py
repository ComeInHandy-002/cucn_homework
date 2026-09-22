"""Generate the classroom demonstration and presentation handbook."""

from pathlib import Path

from docx import Document

from docx_common import (
    add_code,
    add_cover,
    add_intro,
    add_label_paragraph,
    add_list,
    add_table,
    configure_document,
    save,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "学号_姓名_工业安全智能监测系统_展示讲解演示手册.docx"


def build() -> None:
    document = Document()
    configure_document(document, "工业安全智能监测系统展示讲解演示手册")
    add_cover(
        document,
        "基于 YOLOv8 与 ByteTrack 的工业作业人员安全防护与危险区域智能监测系统",
        "工业安全检测课程项目",
        "展示讲解演示手册",
    )

    document.add_heading("1 演示目标", level=1)
    add_intro(
        document,
        "用 8 分钟完成一条可见、可操作、可追溯的安全监测闭环。",
        "演示从系统登录开始，依次展示图片检测、视频任务、危险区域、告警处理、任务结果和模型实验。"
        "现场使用已归档的本地素材，不依赖外网、RTSP、人脸识别或真实控制设备。",
    )
    add_table(
        document,
        ["观众应看到", "对应功能"],
        [
            ["图片保持原比例并显示检测框", "智能检测中心的图片上传与结果预览"],
            ["视频任务有进度和可播放结果", "异步任务、任务中心和结果弹窗"],
            ["违规从发现到关闭有记录", "告警中心、证据图、确认和关闭"],
            ["危险区可以现场绘制", "设备与区域页面的 720 x 405 画布"],
            ["指标可追溯到真实文件", "模型实验页的 E7 E6 对比与数据质量"],
        ],
        widths=[3.1, 3.5],
    )

    document.add_heading("2 演示前检查", level=1)
    add_list(
        document,
        [
            "确认 models/expanded-ppe-yolov8s-hardcase12.pt 存在，文件约 22.5 MB。",
            "执行环境自检，确认 Python 3.12.13、Ultralytics 和 CUDA 状态。",
            "启动 8010 端口并先访问 /health，确认 backend 为 ultralytics、model_available 为 true。",
            "使用 admin/admin123 登录，确认顶部服务状态与 WebSocket 状态正常。",
            "在内置素材下拉框确认 8 张图片与 5 段视频可见。",
            "演示前关闭无关标签页和通知，保留工作台与接口文档。",
        ],
        numbered=True,
    )
    add_code(
        document,
        ".\\.venv\\Scripts\\python.exe scripts\\check_environment.py\n"
        ".\\.venv\\Scripts\\python.exe -m uvicorn app.main:app --host 127.0.0.1 --port 8010",
    )
    add_label_paragraph(document, "访问地址：", "http://127.0.0.1:8010/")

    document.add_heading("3 八分钟演示路线", level=1)
    add_table(
        document,
        ["时间", "操作", "画面证据", "讲解重点"],
        [
            ["0:00-0:40", "登录并看总览", "四项指标、服务状态、WebSocket", "系统已连接真实 E7 模型和事件通道"],
            ["0:40-2:10", "载入图片样例并检测", "标注预览、分类计数、事件判定", "检测结果以画面和表格展示，JSON 只做追溯"],
            ["2:10-3:40", "提交视频样例", "进度条、任务 ID、结果按钮", "视频在后台运行，UI 不阻塞"],
            ["3:40-4:50", "新增监控点并画危险区", "设备卡片、多边形和坐标", "脚底中心点进入多边形后触发入侵规则"],
            ["4:50-5:50", "处理告警", "证据图、状态筛选、确认和关闭", "事件具备时间、轨迹、监控点和处理状态"],
            ["5:50-6:50", "查看任务结果", "可播放媒体、目标与事件统计", "结果弹窗优先展示媒体，不把原始 JSON 当主画面"],
            ["6:50-8:00", "打开模型实验并收束", "E7 E6 对比、分项指标、数据划分", "说明实际提升与代价，并主动交代实验边界"],
        ],
        widths=[0.85, 1.45, 2.05, 2.25],
        center_columns={0},
        font_size=8.5,
    )

    document.add_heading("4 图片检测演示", level=1)
    document.add_heading("4.1 推荐素材", level=2)
    add_table(
        document,
        ["素材", "适合展示", "预期讲解"],
        [
            ["02 多人施工 PPE识别", "多人和 PPE 框", "展示 person helmet vest 的计数和关联"],
            ["05 缺少安全帽 有背心", "典型违规", "展示 no_helmet 事件和证据"],
            ["08 单人PPE参考", "合规场景", "说明合规样本不会强行生成告警"],
        ],
        widths=[2.2, 1.6, 2.8],
    )

    document.add_heading("4.2 操作与讲解", level=2)
    add_list(
        document,
        [
            "进入智能检测，选择内置样例并点击载入图片样例。预览必须保持原图比例。",
            "点击开始智能检测。等待状态由推理中变为检测完成。",
            "先指向左侧标注画面，再说明右侧人员、安全帽、反光背心数量和事件结论。",
            "展开原始 JSON 只用于说明接口可追溯，随后收起，避免让 JSON 成为演示主体。",
        ],
        numbered=True,
    )
    add_label_paragraph(
        document,
        "讲解示例：",
        "系统先检测人员、安全帽和反光背心，再按人体区域进行 PPE 关联。"
        "单帧图片直接返回结构化结果；视频还会经过轨迹确认、连续帧确认和冷却。",
    )

    document.add_heading("5 视频任务演示", level=1)
    add_table(
        document,
        ["素材", "时长", "用途"],
        [
            ["03 施工人员行走 跟踪", "约 9.97 秒", "单人 PPE 和稳定轨迹，适合快速演示"],
            ["02 施工人员缺安全帽 违规", "约 5.07 秒", "短视频违规流程，分辨率较高"],
            ["04 城市施工人员 PPE", "约 12.97 秒", "多人场景和轨迹过滤"],
        ],
        widths=[2.55, 1.1, 2.95],
        center_columns={1},
    )
    add_list(
        document,
        [
            "选择内置视频，点击载入视频样例，再提交视频任务。",
            "说明任务状态依次为 queued、running、completed，进度由前端自动轮询。",
            "任务完成后点击查看可视化结果，确认弹窗中出现视频、帧数、目标计数和事件统计。",
            "任务中心可重新打开历史结果；查看原始 JSON仅用于技术追溯。",
        ],
        numbered=True,
    )
    document.add_paragraph(
        "视频人员轨迹必须连续命中 3 次才进入输出。PPE 规则使用 15 次人员观测宽限与预热；"
        "违规连续 3 帧后触发，同一人员同一规则 10 秒内只生成一次告警。"
    )

    document.add_heading("6 设备与危险区域演示", level=1)
    add_table(
        document,
        ["字段", "演示值"],
        [
            ["监控点名称", "一号车间入口"],
            ["视频源", "local"],
            ["安装位置", "A 区北侧"],
            ["区域名称", "主危险区"],
            ["示例坐标", "(160,220) (540,220) (610,365) (105,365)"],
        ],
        widths=[1.75, 4.85],
        center_columns={0},
    )
    add_list(
        document,
        [
            "新增监控点，确认设备卡片和图片、视频绑定下拉框同步更新。",
            "选择监控点，在画布按顺序点击至少 3 个顶点。",
            "演示撤销和清空后重新绘制，再点击保存并启用危险区域。",
            "说明规则使用人员框脚底中心点判断入区，边界点按进入处理。",
        ],
        numbered=True,
    )

    document.add_heading("7 告警闭环演示", level=1)
    add_table(
        document,
        ["状态", "含义", "现场动作"],
        [
            ["open", "新事件等待处理", "按事件类型或状态筛选，打开证据图"],
            ["acknowledged", "人工已确认", "点击确认，观察总览待处理数量变化"],
            ["resolved", "事件已关闭", "点击关闭，验证列表和指标同步刷新"],
        ],
        widths=[1.4, 2.05, 3.15],
        center_columns={0},
    )
    add_label_paragraph(
        document,
        "讲解示例：",
        "告警不是一次性弹窗，而是带证据、状态和处理时间的数据库记录。"
        "WebSocket 负责实时提示，REST 接口负责查询和状态更新。",
    )

    document.add_heading("8 模型与实验讲解", level=1)
    add_table(
        document,
        ["指标", "E7", "E6", "答辩表述"],
        [
            ["Precision", "0.7977", "0.8332", "精确率下降"],
            ["Recall", "0.6620", "0.6441", "召回率提升"],
            ["mAP@0.5", "0.7200", "0.7151", "低 IoU 综合指标提升"],
            ["mAP@0.5:0.95", "0.4441", "0.4538", "高 IoU 指标下降"],
            ["helmet Recall", "0.4548", "0.4353", "安全帽漏检有所改善"],
        ],
        widths=[1.4, 1.05, 1.05, 3.1],
        center_columns={0, 1, 2},
    )
    document.add_paragraph(
        "推荐表述：E7 选择了更偏召回的部署取舍，改善了安全帽召回和 mAP@0.5，"
        "但 Precision 与 mAP@0.5:0.95 有小幅下降，因此不能表述为所有指标全面提升。"
    )
    add_table(
        document,
        ["数据事实", "数值"],
        [
            ["基础合并集", "train 8,276，val 2,184，test 1,163；共 11,623 张图片"],
            ["基础标注框", "train 20,695，val 5,111，test 3,047；共 28,853 个框"],
            ["E7 训练补充", "3 张原图加 180 张相关增强，共 1,037 个框"],
            ["E7 有效训练输入", "8,459 张图片，21,732 个框"],
        ],
        widths=[2.0, 4.6],
    )

    document.add_heading("9 常见答辩问题", level=1)
    add_table(
        document,
        ["问题", "建议回答"],
        [
            ["为什么使用 YOLOv8s", "在本机 GPU 上兼顾精度和演示速度；YOLOv8n 保留为轻量对照。"],
            ["ByteTrack 在系统中做什么", "为视频人员框维持 track_id，支持连续帧确认和按人员去重。"],
            ["PPE 如何关联到人员", "安全帽或背心中心需落入人员框，并满足人体纵向区域限制；正向 PPE 证据优先处理冲突。"],
            ["为何还有误检漏检", "复杂遮挡、小目标、颜色和视角仍会影响检测；E7 改善召回但没有消除所有误差。"],
            ["hard-case 是否属于独立测试", "不属于。它参与了定向训练，只用于防止已知问题回归。"],
            ["是否达到 15 FPS", "尚未完成标准化 720p 端到端基准，不使用目标值代替实测值。"],
            ["为何暂不做 RTSP 和人脸识别", "本阶段范围限定为图片、MP4 和本地摄像头；身份识别不属于安全防护规则。"],
        ],
        widths=[2.0, 4.6],
        font_size=8.8,
    )

    document.add_heading("10 现场故障处理", level=1)
    add_table(
        document,
        ["问题", "立即处理"],
        [
            ["页面打不开", "先访问 /health；若端口冲突，换端口启动并使用新地址。"],
            ["登录失败", "确认服务已初始化数据库，使用 admin/admin123；不要刷新数据库文件。"],
            ["图片无标注结果", "检查模型状态、文件格式和浏览器缓存；结果区应显示错误信息。"],
            ["视频长时间 running", "先演示其他页面，稍后在任务中心刷新；不要重复提交同一大视频。"],
            ["结果弹窗无媒体", "检查 results 静态路径与任务摘要；原始 JSON 只作为备用追溯。"],
            ["现场模型不可用", "可展示既有结果和报告，但必须明确这是历史实测证据；fallback 仅能演示接口流程。"],
        ],
        widths=[2.0, 4.6],
    )

    document.add_heading("11 最终演示检查表", level=1)
    add_table(
        document,
        ["序号", "检查项", "结果"],
        [
            ["1", "首页服务状态和 WebSocket 在线", "□"],
            ["2", "登录遮罩消失，显示 admin 账户", "□"],
            ["3", "图片预览保持比例，检测框和统计同时出现", "□"],
            ["4", "视频任务完成，结果弹窗可播放媒体", "□"],
            ["5", "监控点创建和危险区保存成功", "□"],
            ["6", "告警证据可打开，确认和关闭状态生效", "□"],
            ["7", "任务中心可回看图片与视频历史结果", "□"],
            ["8", "模型实验显示 E7 E6 指标与 train val test 数据", "□"],
            ["9", "桌面和移动宽度下无文字或控件重叠", "□"],
            ["10", "主动说明未完成的 80 epoch、独立自建集、15 FPS 与消融", "□"],
        ],
        widths=[0.7, 5.15, 0.75],
        center_columns={0, 2},
    )

    document.add_heading("12 一分钟收束讲稿", level=1)
    document.add_paragraph(
        "本项目已经形成从目标检测、人员跟踪、PPE 关联、危险区域规则到告警闭环的完整工程链路。"
        "当前默认 E7 模型在独立 test 上的 Recall 为 0.6620、mAP@0.5 为 0.7200，"
        "安全帽 Recall 相比 E6 从 0.4353 提升到 0.4548；同时 Precision 和 mAP@0.5:0.95 有小幅下降。"
        "系统通过 3 次轨迹命中、3 帧事件确认和 10 秒冷却减少视频瞬时误报，"
        "并用任务、证据截图和事件状态支持结果追溯。当前成果可以完成课程演示和阶段提交，"
        "但 80 epoch、独立自建场景测试、标准化 720p 性能验收和定量消融仍应继续完成。"
    )

    save(document, OUTPUT)


if __name__ == "__main__":
    build()
