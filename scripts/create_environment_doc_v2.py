"""生成《工业安全智能监测系统 环境搭建与运行说明》（实名版，双服务现状）。

覆盖当前系统全部组成部分：Python 算法域、Spring Boot 管理域、Vue3 前端、
MySQL/PostgreSQL、Docker 可选部署。封面学生信息来自 results/midterm/student.json。
"""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document

from docx_common import (
    add_code,
    add_cover,
    add_intro,
    add_label_paragraph,
    add_list,
    add_table,
    configure_document,
    save,
)

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Z202508090311_朱文杰_工业安全智能监测系统_环境搭建与运行说明.docx"

try:
    STUDENT = json.loads((ROOT / "results/midterm/student.json").read_text(encoding="utf-8"))
except FileNotFoundError:
    STUDENT = {}


def build() -> None:
    document = Document()
    configure_document(document, "工业安全智能监测系统环境搭建与运行说明")
    add_cover(
        document,
        "基于 YOLOv8 与 ByteTrack 的工业作业人员安全防护与危险区域智能监测系统",
        f"{STUDENT.get('学号', '学号')} {STUDENT.get('姓名', '姓名')}",
        "环境搭建与运行说明",
    )

    document.add_heading("1 文档用途与系统组成", level=1)
    add_intro(
        document,
        "本说明用于在 Windows 11 本机从零复现当前系统的完整运行。",
        "系统为前后端分离的双服务架构：Python 算法域（FastAPI :8010，YOLOv8 推理/跟踪/规则/WebSocket）"
        "与 Java 管理域（Spring Boot :8081，用户管理、数据集与模型任务，MySQL 库 safety_admin），"
        "前端为 Vue3 工程（构建产物由算法域伺服）。全部命令均在本机实测通过。",
    )
    add_table(
        document,
        ["组成部分", "技术栈", "默认端口"],
        [
            ["算法域 API + 工作台", "Python 3.12 + FastAPI + Ultralytics", "8010"],
            ["管理域 API", "JDK 24 + Spring Boot 3.4 + JPA", "8081"],
            ["业务数据库", "MySQL 8.0（库 safety_admin）", "3306"],
            ["前端构建（一次性）", "Node.js 22 + Vite + Vue3", "—"],
        ],
    )

    document.add_heading("2 环境要求", level=1)
    document.add_heading("2.1 硬件与基础软件", level=2)
    add_table(
        document,
        ["项目", "要求", "说明"],
        [
            ["GPU（可选）", "NVIDIA GPU 8GB + CUDA 12.6 驱动", "无 GPU 时自动回退 CPU，功能不受影响仅速度下降"],
            ["内存/磁盘", "16GB 内存 / 20GB 可用磁盘", "权重与数据集占用大头"],
            ["Python", "3.12.x", "算法域运行时"],
            ["JDK", "OpenJDK 24（bytecode 21）", "管理域运行时"],
            ["MySQL", "8.0.x", "管理域业务库"],
            ["Node.js / npm", "22.x / 10.x", "仅前端构建需要，运行时不依赖"],
            ["Maven", "3.9.x", "管理域构建（可用自带命令路径）"],
            ["Git", "2.5x", "获取源码"],
        ],
    )
    document.add_heading("2.2 安装验证", level=2)
    add_intro(
        document,
        "安装完成后逐项执行以下命令，输出应与示例一致（真实运行记录）：",
        "",
    )
    add_code(
        document,
        "python --version      # Python 3.12.13\n"
        "java -version         # openjdk version \"24.0.2\"\n"
        "git --version         # git version 2.51.0.windows.1\n"
        "mysql --version       # MySQL 8.0.43\n"
        "node --version        # v22.23.2\n"
        "docker --version      # Docker version 29.7.2",
    )

    document.add_heading("3 安装步骤", level=1)
    document.add_heading("3.1 获取源码", level=2)
    add_code(document, "git clone https://github.com/ComeInHandy-002/cucn_homework.git\ncd cucn_homework")
    document.add_heading("3.2 Python 算法域", level=2)
    add_label_paragraph(document, "步骤 1：", "创建虚拟环境并安装依赖（requirements.txt 含 FastAPI/Ultralytics/SQLAlchemy 等）。")
    add_code(document, "uv venv\nuv pip install -r requirements.txt\n# 无 uv 时: python -m venv .venv && .venv\\Scripts\\activate && pip install -r requirements.txt")
    add_label_paragraph(document, "步骤 2：", "安装 PyTorch（有 NVIDIA GPU 用 CUDA 版，否则 CPU 版）。")
    add_code(
        document,
        "# CUDA 12.x\npip install torch torchvision --index-url https://download.pytorch.org/whl/cu126\n# 仅 CPU\npip install torch torchvision --index-url https://download.pytorch.org/whl/cpu",
    )
    add_label_paragraph(document, "步骤 3：", "放置模型权重（E7 部署权重，提交包内已含，或向作者索取）：")
    add_code(document, "models/expanded-ppe-yolov8s-hardcase12.pt   # 约 22MB")
    document.add_heading("3.3 MySQL 管理域数据库", level=2)
    add_intro(document, "启动 MySQL 服务后创建业务库（表结构由 JPA 首次启动自动创建）：", "")
    add_code(document, "mysql -uroot -p -e \"CREATE DATABASE IF NOT EXISTS safety_admin DEFAULT CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci;\"")
    document.add_heading("3.4 Spring Boot 管理域", level=2)
    add_intro(document, "Maven 依赖经阿里云镜像拉取（仓库内 .mvn/settings-mirror.xml 已配置）。设置环境变量后启动：", "")
    add_code(
        document,
        "cd backend-java\n"
        "set MYSQL_PASSWORD=你的数据库密码\n"
        "set MYSQL_HOST=127.0.0.1\n"
        "set MYSQL_DB=safety_admin\n"
        "\"D:\\Tools\\maven\\apache-maven-3.9.9\\bin\\mvn.cmd\" -s .mvn/settings-mirror.xml spring-boot:run",
    )
    add_label_paragraph(document, "说明：", "首次启动自动建 users/datasets/model_train_tasks 三表并写入内置管理员 admin/admin123（BCrypt 加密）。")
    document.add_heading("3.5 前端构建（可选，仓库已含构建产物）", level=2)
    add_code(document, "cd frontend\nnpm install\nnpm run build   # 产物输出到 app/static/dist/，算法域直接伺服")

    document.add_heading("4 启动与验证", level=1)
    document.add_heading("4.1 启动顺序", level=2)
    add_table(
        document,
        ["顺序", "服务", "命令（仓库根目录执行）"],
        [
            ["1", "MySQL", "系统服务或 mysqld --console"],
            ["2", "管理域 Spring Boot", "见 3.4（约 20 秒就绪）"],
            ["3", "算法域 FastAPI", ".venv\\Scripts\\python.exe -m uvicorn app.main:app --host 127.0.0.1 --port 8010"],
        ],
    )
    document.add_heading("4.2 健康检查", level=2)
    add_code(
        document,
        "curl http://127.0.0.1:8010/health      # model_available=true 即权重加载成功\n"
        "curl http://127.0.0.1:8081/            # 返回管理域引导页 HTML",
    )
    document.add_heading("4.3 访问入口与账号", level=2)
    add_table(
        document,
        ["入口", "地址", "账号"],
        [
            ["监测工作台（主入口）", "http://127.0.0.1:8010", "admin / admin123"],
            ["用户管理（管理域）", "http://127.0.0.1:8010/admin/users", "admin / admin123（独立登录）"],
            ["数据集与模型管理", "http://127.0.0.1:8010/admin/models", "同上"],
            ["管理域 API 引导页", "http://127.0.0.1:8081/", "无页面，仅提示"],
        ],
    )

    document.add_heading("5 Docker 方式（可选）", level=1)
    add_intro(document, "整体容器化：nginx（前端+反代，公开 8000）+ api（FastAPI）+ db（PostgreSQL 16）。", "")
    add_code(
        document,
        "docker compose up --build -d\n# NVIDIA 主机启用 GPU 推理：\ndocker compose -f docker-compose.yml -f docker-compose.gpu.yml up --build",
    )
    add_label_paragraph(document, "镜像拉取受限时：", "先用镜像源补齐基础镜像再 compose up（仓库 README 有可复制脚本）。")

    document.add_heading("6 自动化测试", level=1)
    add_intro(document, "算法域测试套件 65 个用例（检测几何过滤/跟踪/规则/API 契约/视频流水线），管理域提供全接口测试矩阵：", "")
    add_code(document, ".venv\\Scripts\\python.exe -m pytest -q\n# 预期输出: 65 passed")

    document.add_heading("7 常见问题排查", level=1)
    add_list(
        document,
        [
            "GitHub 克隆/推送超时：本机代理环境下执行 git -c http.proxy=http://127.0.0.1:7897 push（端口按实际代理调整），或在 ~/.gitconfig 按域名配置代理。",
            "Maven 依赖拉取失败：确认使用 -s .mvn/settings-mirror.xml（阿里云镜像）；基础 Docker 镜像同理可用 DaoCloud 源。",
            "8000/8010 端口被占用：netstat -ano | findstr :8010 找到进程后 taskkill /PID <pid> /F，或改用其他端口启动。",
            "/health 中 model_available=false：检查 models/ 下权重文件名与 .env 的 MODEL_PATH 一致；无权重时系统进入 fallback 演示模式，接口可用但检测为合成结果。",
            "管理域接口全部 401：未携带令牌或令牌过期，重新登录获取；管理域与监测域令牌相互独立。",
            "浏览器页面空白：前端更新后旧缓存引用失效资源，Ctrl+F5 强制刷新；未知路径已自动重定向回总览。",
            "无 GPU 环境：MODEL_DEVICE 留空自动选择 CPU，视频 batch 可调小（VIDEO_BATCH_SIZE=4）。",
        ],
        numbered=True,
    )

    save(document, OUTPUT)


if __name__ == "__main__":
    build()
