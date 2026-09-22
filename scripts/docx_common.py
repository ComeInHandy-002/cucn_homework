"""Shared Word formatting helpers for the two project handbooks."""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


FONT_CJK = "Microsoft YaHei"
FONT_MONO = "Consolas"
TEXT = "1F2937"
MUTED = "5B6573"
NAVY = "17365D"
PALE_BLUE = "EEF4FA"
BORDER = "D9D9D9"
CODE_BG = "F4F6F8"


def _set_font(run, name: str = FONT_CJK, size: float | None = None, color: str | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _set_style_font(style, name: str, size: float, color: str = TEXT, bold: bool | None = None) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), BORDER)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def _set_footer(section, short_title: str) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    run = paragraph.add_run(f"{short_title}  |  ")
    _set_font(run, size=8.5, color=MUTED)
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    paragraph._p.append(page)


def configure_document(document: Document, short_title: str) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    _set_footer(section, short_title)

    normal = document.styles["Normal"]
    _set_style_font(normal, FONT_CJK, 10.5, TEXT)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5.5)

    title = document.styles["Title"]
    _set_style_font(title, FONT_CJK, 24, "000000", True)
    title.paragraph_format.space_after = Pt(12)
    title_p_pr = title.element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    heading1 = document.styles["Heading 1"]
    _set_style_font(heading1, FONT_CJK, 15, "000000", True)
    heading1.paragraph_format.space_before = Pt(14)
    heading1.paragraph_format.space_after = Pt(7)
    heading1.paragraph_format.keep_with_next = True

    heading2 = document.styles["Heading 2"]
    _set_style_font(heading2, FONT_CJK, 12, "000000", True)
    heading2.paragraph_format.space_before = Pt(9)
    heading2.paragraph_format.space_after = Pt(4)
    heading2.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = document.styles[style_name]
        _set_style_font(style, FONT_CJK, 10.5, TEXT)
        style.paragraph_format.space_after = Pt(3)

    document.core_properties.title = short_title
    document.core_properties.subject = "工业安全智能监测系统课程项目交付文档"
    document.core_properties.author = "学生"


def add_cover(document: Document, title: str, subtitle: str, document_name: str) -> None:
    top = document.add_paragraph()
    top.paragraph_format.space_before = Pt(64)
    top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = top.add_run("课程项目阶段交付")
    _set_font(run, size=11, color=MUTED)

    heading = document.add_paragraph(style="Title")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.add_run(title)
    heading_p_pr = heading._p.get_or_add_pPr()
    heading_border = heading_p_pr.find(qn("w:pBdr"))
    if heading_border is not None:
        heading_p_pr.remove(heading_border)

    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(30)
    run = sub.add_run(subtitle)
    _set_font(run, size=12, color=MUTED)

    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run(document_name)
    _set_font(run, size=17, color=NAVY)
    run.bold = True

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(55)
    meta.paragraph_format.line_spacing = 1.8
    run = meta.add_run(
        f"学号  ____________________\n姓名  ____________________\n提交日期  {date.today().isoformat()}"
    )
    _set_font(run, size=11, color=TEXT)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_intro(document: Document, lead: str, details: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    lead_run = paragraph.add_run(lead)
    _set_font(lead_run, size=11)
    lead_run.bold = True
    detail_run = paragraph.add_run(details)
    _set_font(detail_run, size=10.5)


def add_table(
    document: Document,
    headers: Sequence[str],
    rows: Iterable[Sequence[str]],
    widths: Sequence[float] | None = None,
    center_columns: set[int] | None = None,
    font_size: float = 9.2,
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    _repeat_header(table.rows[0])
    centers = center_columns or set()

    for column, (cell, value) in enumerate(zip(table.rows[0].cells, headers)):
        _set_cell_shading(cell, NAVY)
        _set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.1
        run = paragraph.add_run(str(value))
        _set_font(run, size=font_size, color="FFFFFF")
        run.bold = True
        if widths:
            cell.width = Inches(widths[column])

    for row_index, values in enumerate(rows):
        row = table.add_row()
        _prevent_row_split(row)
        for column, (cell, value) in enumerate(zip(row.cells, values)):
            if row_index % 2:
                _set_cell_shading(cell, PALE_BLUE)
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column in centers else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            run = paragraph.add_run(str(value))
            _set_font(run, size=font_size)
            if widths:
                cell.width = Inches(widths[column])
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_code(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), CODE_BG)
    p_pr.append(shd)
    run = paragraph.add_run(text)
    _set_font(run, FONT_MONO, 8.7, TEXT)


def add_list(document: Document, items: Sequence[str], numbered: bool = False) -> None:
    for index, item in enumerate(items, start=1):
        if numbered:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.28)
            paragraph.paragraph_format.first_line_indent = Inches(-0.28)
            paragraph.paragraph_format.space_after = Pt(3)
            value = f"{index}.  {item}"
        else:
            paragraph = document.add_paragraph(style="List Bullet")
            value = item
        run = paragraph.add_run(value)
        _set_font(run, size=10.5)


def add_label_paragraph(document: Document, label: str, text: str) -> None:
    paragraph = document.add_paragraph()
    first = paragraph.add_run(label)
    _set_font(first, size=10.5)
    first.bold = True
    second = paragraph.add_run(text)
    _set_font(second, size=10.5)


def save(document: Document, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(output)
