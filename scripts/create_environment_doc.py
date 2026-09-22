"""Generate the verified environment and run guide for the project."""

from pathlib import Path

from docx import Document

from docx_common import (
    add_code,
    add_cover,
    add_intro,
    add_label_paragraph,
    add_list,
    add_table,
    configure_document,
    save,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "学号_姓名_工业安全智能监测系统_环境搭建与运行说明.docx"


def build() -> None:
    document = Document()
    configure_document(document, "工业安全智能监测系统环境搭建与运行说明")
    add_cover(
        document,
        "基于 YOLOv8 与 ByteTrack 的工业作业人员安全防护与危险区域智能监测系统",
        "工业安全检测课程项目",
        "环境搭建与运行说明",
    )

    document.add_heading("1 文档用途与当前状态", level=1)
    add_intro(
        document,
        "本说明用于在 Windows 11 本机复现当前项目。",
        "内容覆盖运行环境、数据与模型、启动方法、接口校验、自动化测试和故障排查。"
        "当前默认模型为 E7 12 epoch 定向精调权重，服务、工作台、数据集和回归报告均已落盘。",
    )
    add_table(
        document,
        ["交付项", "当前状态", "可核验证据"],
        [
            ["后端与工作台", "FastAPI 与六页单页工作台可运行", "app 目录与 49 项自动化测试"],
            ["当前模型", "E7 hard-case 12 epoch，best epoch 7", "models/expanded-ppe-yolov8s-hardcase12.pt"],
            ["数据", "合并集 11,623 张图，28,853 个框", "expanded_ppe_kaggle 数据清单与校验报告"],
            ["本机算力", "RTX 4070 Laptop GPU，CUDA 12.6 可用", "scripts/check_environment.py"],
            ["定向回归", "3 张图片与 2 段视频门禁通过", "results 下 hardcase12 回归报告"],
        ],
        widths=[1.35, 2.35, 2.9],
        center_columns={0},
    )

    document.add_heading("2 系统组成", level=1)
    document.add_paragraph(
        "浏览器工作台调用 FastAPI 接口。图片任务同步返回结果，视频任务进入单 Worker 队列；"
        "推理链路依次完成 YOLOv8 检测、人员轨迹确认、PPE 关联和危险区域规则判断，"
        "再把任务、事件、证据图和结果媒体写入数据库与本地目录。"
    )
    add_code(
        document,
        "Browser UI\n"
        "  -> FastAPI API / JWT / WebSocket\n"
        "      -> Image inference or video job queue\n"
        "          -> YOLOv8s -> track confirmation -> PPE association -> zone rules\n"
        "              -> SQLite or PostgreSQL + results/ + uploads/",
    )
    add_table(
        document,
        ["层次", "实现", "职责"],
        [
            ["展示层", "HTML CSS JavaScript", "登录、总览、检测、任务、实验、告警、设备与区域"],
            ["接口层", "FastAPI JWT WebSocket", "鉴权、上传、任务查询、事件处理、实时推送"],
            ["推理层", "Ultralytics YOLOv8s", "检测 person helmet vest 并输出框与置信度"],
            ["时序层", "ByteTrack 思路与规则引擎", "3 次命中确认、PPE 关联、3 帧事件确认、10 秒冷却"],
            ["数据层", "SQLite 或 PostgreSQL", "用户、摄像头、危险区、任务和安全事件"],
            ["文件层", "Docker volume 或本地目录", "权重、上传文件、预览、证据截图和结果媒体"],
        ],
        widths=[1.0, 2.0, 3.6],
        center_columns={0},
    )

    document.add_heading("3 已验证运行环境", level=1)
    add_table(
        document,
        ["类别", "实测配置"],
        [
            ["操作系统", "Windows 11"],
            ["Python", "3.12.13，项目由 .python-version 固定"],
            ["PyTorch", "2.6.0+cu126，TorchVision 0.21.0+cu126"],
            ["GPU", "NVIDIA GeForce RTX 4070 Laptop GPU，CUDA 12.6 可用"],
            ["Web", "FastAPI 0.141.1，Uvicorn 0.53.0，SQLAlchemy 2.0.53，Pydantic 2.13.5"],
            ["视觉", "Ultralytics 8.4.152，OpenCV 5.0.0，Pillow 12.3.0，NumPy 2.5.3"],
            ["数据库", "开发环境默认 SQLite；Docker Compose 配置 PostgreSQL 16"],
            ["模型输入", "640 px，confidence 0.25，设备留空时由 Ultralytics 自动选择 CUDA 或 CPU"],
        ],
        widths=[1.45, 5.15],
        center_columns={0},
    )

    document.add_heading("4 目录与关键文件", level=1)
    add_table(
        document,
        ["路径", "说明"],
        [
            ["app", "FastAPI、数据库、检测器、跟踪器、规则引擎和静态工作台"],
            ["data", "公开数据、Kaggle 合并集、hard-case、演示素材和来源清单"],
            ["models", "当前模型、历史权重和 YOLOv8n YOLOv8s 对照权重"],
            ["scripts", "安装、下载、校验、训练、评估、回归与文档生成脚本"],
            ["tests", "检测、跟踪、规则、API、WebSocket、前端与回归工具测试"],
            ["results", "指标、校验报告、任务 JSON、预览、证据图和结果视频"],
            ["uploads", "用户上传的图片与视频，运行时保留"],
            ["docs", "系统设计、API、数据、实验、测试、用户手册与答辩材料"],
        ],
        widths=[1.35, 5.25],
        center_columns={0},
    )

    document.add_heading("5 本机安装与启动", level=1)
    document.add_heading("5.1 一键安装", level=2)
    document.add_paragraph("在项目根目录打开 PowerShell。首次安装 CUDA 12.6 版本依赖时执行：")
    add_code(document, ".\\scripts\\setup.ps1 -Gpu")
    document.add_paragraph("依赖已安装时只做环境检查：")
    add_code(document, ".\\scripts\\setup.ps1 -SkipInstall")
    add_label_paragraph(
        document,
        "安装路径说明：",
        "项目虚拟环境、uv 缓存和构建产物保存在项目所在的 D 盘目录；脚本不会主动把大体积依赖写入其他盘。",
    )

    document.add_heading("5.2 环境变量", level=2)
    document.add_paragraph("首次运行会从 .env.example 生成 .env。当前关键配置如下：")
    add_code(
        document,
        "MODEL_PATH=models/expanded-ppe-yolov8s-hardcase12.pt\n"
        "MODEL_CONFIDENCE=0.25\n"
        "MODEL_IMAGE_SIZE=640\n"
        "MODEL_DEVICE=\n"
        "QUEUE_WORKERS=1\n"
        "MAX_UPLOAD_MB=100",
    )
    document.add_paragraph(
        "MODEL_DEVICE 留空表示自动选择设备；需要固定第一块 GPU 时可设置为 0。"
        "生产环境必须替换 SECRET_KEY 和默认账户密码，并限制 CORS_ORIGINS。"
    )

    document.add_heading("5.3 启动服务", level=2)
    add_code(
        document,
        ".\\.venv\\Scripts\\python.exe -m uvicorn app.main:app --host 127.0.0.1 --port 8010",
    )
    add_list(
        document,
        [
            "工作台地址：http://127.0.0.1:8010/",
            "接口文档：http://127.0.0.1:8010/docs",
            "健康检查：http://127.0.0.1:8010/health",
            "开发账号：admin，密码：admin123。",
        ],
    )

    document.add_heading("5.4 Docker Compose", level=2)
    add_code(document, "docker compose up --build")
    document.add_paragraph(
        "Compose 使用 PostgreSQL 16，并挂载 models、uploads、results 和 data。"
        "容器配置保持 CPU 安全模式；如需容器内 CUDA，必须改用 CUDA 基础镜像并配置 NVIDIA Container Toolkit。"
    )

    document.add_heading("6 数据与模型", level=1)
    add_table(
        document,
        ["划分", "图片", "标注框", "用途"],
        [
            ["train", "8,276", "20,695", "E6 基础训练集"],
            ["val", "2,184", "5,111", "训练选择与调参"],
            ["test", "1,163", "3,047", "独立指标，未挂载 hard-case"],
            ["hard-case train", "3 原图加 180 增强", "1,037", "只用于已知问题定向修复"],
            ["E7 有效 train", "8,459", "21,732", "E6 train 加 hard-case"],
        ],
        widths=[1.35, 1.55, 1.35, 2.35],
        center_columns={0, 1, 2},
    )
    document.add_paragraph(
        "基础合并集来自 Construction-PPE、SH17 和 Kaggle PPE Detection，按图片哈希去重并保留来源、许可证与校验记录。"
        "3 张 hard-case 原图及其 180 张相关增强不是 183 个独立场景，也不能替代计划中的独立自建测试集。"
    )

    document.add_heading("6.1 E7 独立测试结果", level=2)
    add_table(
        document,
        ["指标", "E7", "E6", "变化"],
        [
            ["Precision", "0.7977", "0.8332", "下降 0.0355"],
            ["Recall", "0.6620", "0.6441", "提升 0.0179"],
            ["F1", "0.7235", "0.7266", "下降 0.0031"],
            ["mAP@0.5", "0.7200", "0.7151", "提升 0.0049"],
            ["mAP@0.5:0.95", "0.4441", "0.4538", "下降 0.0097"],
            ["helmet Recall", "0.4548", "0.4353", "提升 0.0195"],
        ],
        widths=[1.75, 1.25, 1.25, 2.35],
        center_columns={0, 1, 2, 3},
    )
    document.add_paragraph(
        "E7 提高了 Recall、mAP@0.5 和安全帽召回率，同时牺牲了部分 Precision 与高 IoU 指标。"
        "当前部署选择 E7 是为了减少漏检，但论文和答辩必须保留这一真实取舍。"
    )

    document.add_heading("6.2 定向回归边界", level=2)
    add_list(
        document,
        [
            "三人施工图：3 person、3 helmet、2 vest，只生成 no_vest，机械区域不再误检为人员。",
            "黑帽工厂图：2 person、2 helmet、2 vest，不生成 no_helmet。",
            "施工通道图：1 person、1 helmet、1 vest，无违规事件。",
            "两段视频分别确认 1 条与 3 条人员轨迹，短暂机械候选不会进入输出、统计或规则。",
        ],
    )
    document.add_paragraph("这些结果只证明指定已知案例通过回归门禁，不属于独立泛化指标。")

    document.add_heading("7 运行验证", level=1)
    document.add_heading("7.1 自动化检查", level=2)
    add_code(
        document,
        ".\\.venv\\Scripts\\python.exe -m pytest -q\n"
        "node --check app\\static\\app.js\n"
        ".\\.venv\\Scripts\\python.exe -m compileall -q app scripts tests\n"
        ".\\.venv\\Scripts\\python.exe scripts\\check_environment.py",
    )
    document.add_paragraph(
        "2026-09-20 复核结果为 49 passed、2 条既有依赖弃用 warning；JavaScript 语法、Python compileall 和环境自检均通过。"
    )

    document.add_heading("7.2 接口验收", level=2)
    add_table(
        document,
        ["操作", "接口", "通过标准"],
        [
            ["健康检查", "GET /health", "backend 为 ultralytics，model_available 为 true，权重为 E7"],
            ["登录", "POST /api/v1/auth/login", "返回 JWT 与 admin 用户信息"],
            ["图片检测", "POST /api/v1/inference/images", "返回检测、事件、JSON 路径和标注预览"],
            ["视频任务", "POST /api/v1/inference/videos", "返回 202，任务最终 completed 或明确失败原因"],
            ["结果摘要", "GET /api/v1/jobs/{id}/result", "返回可视化媒体、分类统计和事件摘要"],
            ["实验页", "GET /api/v1/experiments/summary", "包含 E7 当前模型与 E6 历史对照"],
            ["实时事件", "WS /api/v1/ws/events", "JWT 鉴权、事件推送和 ping pong 正常"],
        ],
        widths=[1.15, 2.2, 3.25],
        center_columns={0},
        font_size=8.8,
    )
    document.add_paragraph(
        "最终 API 实测中，hard_02.jpg 返回 2 person、2 helmet、2 vest 且无事件；"
        "239 帧视频任务完成并生成可播放 WebM，结果共 742 个确认检测且无事件。"
    )

    document.add_heading("8 常见问题", level=1)
    add_table(
        document,
        ["现象", "处理方法"],
        [
            ["8010 端口被占用", "改用其他端口启动，并访问对应 URL。"],
            ["健康检查显示模型不可用", "核对 MODEL_PATH、权重文件和 Ultralytics 安装；不要用 fallback 指标冒充模型结果。"],
            ["CUDA 不可用", "运行 setup.ps1 -Gpu，确认 NVIDIA 驱动与 torch 2.6.0+cu126；也可临时使用 CPU。"],
            ["视频任务失败", "在任务中心查看 error_message，并确认 OpenCV、编码器与文件格式。"],
            ["前端仍显示旧资源", "强制刷新页面；静态文件 URL 已带版本参数。"],
            ["Docker 无法连接数据库", "先确认 Docker 引擎正常，再检查 5432 端口和 postgres_data volume。"],
        ],
        widths=[2.0, 4.6],
    )

    document.add_heading("9 当前实验与部署边界", level=1)
    add_intro(
        document,
        "当前可提交的是完整可运行工程与已复现实验。",
        "以下项目尚未执行，材料中不得写成已完成：80 epoch 长周期训练、300 到 500 帧独立自建场景测试、"
        "720p 端到端不低于 15 FPS 的正式验收，以及无跟踪、仅跟踪、完整时间过滤三组定量消融。",
    )
    add_label_paragraph(
        document,
        "生产部署提醒：",
        "开发版 /results 静态媒体无需登录即可访问，课程本机演示可以使用；公网部署前应改为鉴权下载或短时签名 URL。",
    )
    add_label_paragraph(
        document,
        "结果追溯：",
        "模型指标见 results/expanded-ppe-yolov8s-hardcase12-test-metrics.json；"
        "图片与视频门禁分别见 results/hardcase-regression-hardcase12.json 和 results/video-regression-hardcase12.json。",
    )

    document.add_heading("9.1 交付文件索引", level=2)
    add_table(
        document,
        ["文件或目录", "用途"],
        [
            ["README.md 与 docs/INDEX.md", "项目入口、运行方式和全部文档导航"],
            [".env.example 与 docker-compose.yml", "本机和容器配置，默认均指向 E7 权重"],
            ["models/expanded-ppe-yolov8s-hardcase12.pt", "当前演示部署权重"],
            ["results/*hardcase12*.json", "独立 test、图片门禁和视频门禁证据"],
            ["docs/TEST_REPORT.md", "自动化测试、API 实测、模型结果和剩余边界"],
            ["两份阶段 DOCX", "环境搭建运行说明与展示讲解演示手册"],
        ],
        widths=[3.05, 3.55],
    )
    add_label_paragraph(
        document,
        "最小复现顺序：",
        "运行 setup.ps1 -SkipInstall，启动 8010 端口，访问 /health，登录工作台，再依次执行图片与视频样例。",
    )

    save(document, OUTPUT)


if __name__ == "__main__":
    build()
